"""
Utilidades compartidas para generar los documentos legales de Chekdor.

Diseño: un solo lugar define tipografía, márgenes, numeración y el
  mini-lenguaje de marcado en línea (**negritas**, __cursiva__), para que
  los tres documentos salgan consistentes y sean fáciles de mantener.

Ejecuta este archivo directamente para una prueba de humo.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

# ---------------------------------------------------------------------
# Identidad de marca y tipografía
# ---------------------------------------------------------------------
FUENTE = "Calibri"
AZUL = RGBColor(0x1D, 0x4E, 0xD8)  # azul de marca (themeColor #1d4ed8)
GRIS = RGBColor(0x55, 0x5F, 0x6D)
ROJO = RGBColor(0xB4, 0x23, 0x18)
NEGRO = RGBColor(0x11, 0x11, 0x11)


def documento_nuevo(titulo: str, subtitulo: str | None = None) -> Document:
    """Crea un documento con la configuración de página y estilos de Chekdor."""
    doc = Document()

    # ---------------------------- Página ----------------------------
    for seccion in doc.sections:
        seccion.top_margin = Cm(2.5)
        seccion.bottom_margin = Cm(2.5)
        seccion.left_margin = Cm(2.8)
        seccion.right_margin = Cm(2.5)

    # ---------------------------- Estilos ---------------------------
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = NEGRO
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    # Fuente para caracteres no latinos (acentos y comillas tipográficas
    # se resuelven solos, pero dejarlo explícito evita sorpresas al abrir
    # el archivo en Word de otra configuración regional).
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FUENTE)
    rfonts.set(qn("w:cs"), FUENTE)

    for nombre, tam, color, espacio_antes, espacio_despues in (
        ("Heading 1", 13, AZUL, 16, 6),
        ("Heading 2", 11.5, NEGRO, 12, 4),
        ("Heading 3", 10.5, GRIS, 10, 3),
    ):
        st = doc.styles[nombre]
        st.font.name = FUENTE
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(espacio_antes)
        st.paragraph_format.space_after = Pt(espacio_despues)
        st.paragraph_format.keep_with_next = True

    # --------------------------- Portada ----------------------------
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CHEKDOR")
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = AZUL

    if subtitulo:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitulo)
        r.font.size = Pt(11)
        r.font.color.rgb = GRIS

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo)
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = NEGRO

    doc.add_paragraph()
    return doc


def _escribir_con_marcado(parrafo, texto: str) -> None:
    """
    Mini-lenguaje de marcado en línea:
      **negritas**   __cursiva__   ~~rojo~~

    Suficiente para estos documentos y sin dependencias externas. No anida
    marcas.
    """
    import re

    patron = re.compile(r"(\*\*.+?\*\*|__.+?__|~~.+?~~)", re.DOTALL)
    for trozo in patron.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("__") and trozo.endswith("__"):
            parrafo.add_run(trozo[2:-2]).italic = True
        elif trozo.startswith("~~") and trozo.endswith("~~"):
            run = parrafo.add_run(trozo[2:-2])
            run.bold = True
            run.font.color.rgb = ROJO
        else:
            parrafo.add_run(trozo)


def parrafo(doc, texto: str, estilo: str | None = None, justificar: bool = True):
    p = doc.add_paragraph(style=estilo)
    _escribir_con_marcado(p, texto)
    if justificar and estilo is None:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def vinetas(doc, elementos: list[str]):
    for e in elementos:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        _escribir_con_marcado(p, e)


def numerada(doc, elementos: list[str], inicio: int = 1):
    for i, e in enumerate(elementos, start=inicio):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(f"{i}. ")
        run.bold = True
        _escribir_con_marcado(p, e)


def encabezado_seccion(doc, texto: str, nivel: int = 1):
    """
    Encabezado de cláusula o sección.

    `nivel` 1 usa Heading 1 (azul, 13 pt) y `nivel` 2 usa Heading 2
    (negro, 11.5 pt), para poder tener subtítulos dentro de un anexo sin
    que parezcan una sección al mismo nivel.
    """
    estilo = "Heading 1" if nivel <= 1 else "Heading 2"
    p = doc.add_paragraph(style=estilo)
    p.paragraph_format.keep_with_next = True
    _escribir_con_marcado(p, texto)
    return p


def tabla(doc, encabezados: list[str], filas: list[list[str]], anchos: list[float] | None = None):
    """Tabla con encabezado sombreado y bordes finos."""
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    hdr = t.rows[0].cells
    for i, texto in enumerate(encabezados):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(texto)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(hdr[i], "1D4ED8")

    for fila in filas:
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = ""
            p = celdas[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _escribir_con_marcado(p, str(valor))
            for r in p.runs:
                r.font.size = Pt(9.5)

    if anchos:
        for fila in t.rows:
            for i, ancho in enumerate(anchos):
                fila.cells[i].width = Cm(ancho)

    doc.add_paragraph()
    return t


def _sombrear(celda, color_hex: str):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcpr.append(shd)


def caja_aviso(doc, titulo: str, texto: str, tono: str = "rojo"):
    """
    Bloque destacado de una sola celda. Se usa para las advertencias que el
    usuario NO debe pasar por alto (campos por llenar, lagunas de seguridad).
    """
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    celda = t.rows[0].cells[0]
    celda.text = ""

    relleno = "FEF2F2" if tono == "rojo" else "FFFBEB" if tono == "ambar" else "EFF6FF"
    _sombrear(celda, relleno)

    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = ROJO if tono == "rojo" else RGBColor(0x92, 0x40, 0x0E)

    p2 = celda.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _escribir_con_marcado(p2, texto)
    for r in p2.runs:
        r.font.size = Pt(9.5)

    doc.add_paragraph()
    return t


def firma(doc, etiquetas: list[str]):
    """Bloque de firmas con línea y etiqueta."""
    doc.add_paragraph()
    t = doc.add_table(rows=3, cols=len(etiquetas))
    t.autofit = True
    for i, etiqueta in enumerate(etiquetas):
        p = t.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("_" * 28)
        p2 = t.rows[1].cells[i].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(etiqueta)
        r.bold = True
        r.font.size = Pt(9.5)
        p3 = t.rows[2].cells[i].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run("Nombre y cargo")
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = GRIS
    doc.add_paragraph()
    return t


def _campo(parrafo, instruccion: str):
    """
    Inserta un campo de Word (por ejemplo PAGE o NUMPAGES).

    Un campo no es texto: necesita los elementos fldChar de inicio, la
    instrucción y fldChar de fin. Escribir "PAGE" como texto normal hace
    que Word muestre la palabra, no el número de página.
    """
    run = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruccion} "

    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")

    run._r.append(inicio)
    run._r.append(instr)
    run._r.append(fin)
    return run


def pie_de_pagina(doc, texto: str):
    """
    Pie con el texto a la izquierda y "Página N" alineado a la derecha,
    con el número de página como campo real de Word (no como texto).
    """
    for seccion in doc.sections:
        p = seccion.footer.paragraphs[0]
        p.text = ""
        ancho = section_text_width(seccion)
        p.paragraph_format.tab_stops.add_tab_stop(
            Emu(int(ancho / 2)), WD_TAB_ALIGNMENT.CENTER
        )
        p.paragraph_format.tab_stops.add_tab_stop(Emu(ancho), WD_TAB_ALIGNMENT.RIGHT)

        run = p.add_run(texto)
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS

        # Salta al tabulador derecho y escribe la etiqueta y el número.
        p.add_run("\t\t")
        run = p.add_run("Página ")
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS

        run = _campo(p, "PAGE")
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS


def section_text_width(seccion) -> int:
    """Ancho del área de texto de una sección, en EMU."""
    return int(seccion.page_width - seccion.left_margin - seccion.right_margin)


if __name__ == "__main__":
    d = documento_nuevo("Prueba de humo", "verificación del generador")
    encabezado_seccion(d, "1. Una cláusula de prueba")
    parrafo(d, "Texto con **negritas**, __cursiva__ y ~~rojo de alerta~~.")
    vinetas(d, ["Primera viñeta", "Segunda viñeta"])
    numerada(d, ["Primer inciso", "Segundo inciso"])
    tabla(d, ["A", "B"], [["1", "2"], ["3", "4"]], [6, 6])
    caja_aviso(d, "Aviso", "Esto es una advertencia.", "ambar")
    firma(d, ["Parte A", "Parte B"])
    pie_de_pagina(d, "documento de prueba")
    d.save("_humo.docx")
    print("OK -> _humo.docx")
